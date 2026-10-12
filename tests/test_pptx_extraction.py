"""PPTX extraction tests — decks become one page per slide."""

import pytest

pytest.importorskip("pptx")

from services.document_extractor import DocumentExtractor


@pytest.fixture()
def extractor():
    return DocumentExtractor(enable_ocr=False)


def _build_deck(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()

    # Slide 1: title + body.
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "Quarterly Plan"
    slide1.placeholders[1].text = "Launch the new ingest pipeline."

    # Slide 2: a table plus speaker notes.
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Metrics"
    table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(6), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "System"
    table.cell(0, 1).text = "Uptime"
    table.cell(1, 0).text = "Ingest"
    table.cell(1, 1).text = "99.7%"
    slide2.notes_slide.notes_text_frame.text = "Follow up on the OCR backlog."

    prs.save(str(path))
    return path


def test_pptx_is_supported_extension(extractor):
    assert ".pptx" in extractor.SUPPORTED_EXTENSIONS


def test_slides_become_pages(extractor, tmp_path):
    path = _build_deck(tmp_path / "deck.pptx")
    result = extractor.extract_text(path)
    assert list(result.page_texts) == [1, 2]
    assert "Quarterly Plan" in result.page_texts[1]
    assert "Launch the new ingest pipeline." in result.page_texts[1]
    assert "Metrics" in result.page_texts[2]


def test_table_flattened(extractor, tmp_path):
    path = _build_deck(tmp_path / "deck.pptx")
    result = extractor.extract_text(path)
    assert "System | Uptime" in result.page_texts[2]
    assert "Ingest | 99.7%" in result.page_texts[2]


def test_notes_included(extractor, tmp_path):
    path = _build_deck(tmp_path / "deck.pptx")
    result = extractor.extract_text(path)
    assert "Follow up on the OCR backlog." in result.page_texts[2]


def test_empty_deck_returns_blank_page(extractor, tmp_path):
    from pptx import Presentation
    prs = Presentation()
    prs.save(str(tmp_path / "empty.pptx"))
    result = extractor.extract_text(tmp_path / "empty.pptx")
    assert result.page_texts == {1: ""}
